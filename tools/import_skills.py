#!/usr/bin/env python3
"""스킬/조합 엑셀 → data/skills.json, data/combos.json
사용법: python3 tools/import_skills.py <엑셀 파일>
- '스킬' 시트: id, 카드 이름, 스킬 이름, 종류, 패시브/액티브, 대상, 효과 내용, 수치, 지속 턴, MP 소비, 쿨타임(턴), 해금 레벨, 연출 파일, 비고
- '특정 조합' 시트: 조합 이름, 카드1~4, 효과 대상, 효과 내용, 수치
스킬 효과 구현은 js/skills.js 에서 '스킬 이름'으로 연결됩니다 (효과 내용이 기준).
"""
import json, sys, os
import openpyxl

def num(v, default=0):
    if v is None or v == '' or v == '-': return default
    try: return float(v) if '.' in str(v) else int(v)
    except ValueError: return default

def main(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    root = os.path.join(os.path.dirname(__file__), '..', 'data')
    skills = {}
    fxdir = os.path.join(root, 'fx')
    def fxname(cid, idx, v):
        v = str(v or '').strip()
        if v.endswith('.0'): v = v[:-2]
        if v.isdigit(): v = f'{cid}_{v}'                        # 연출 칸에 스킬 번호만 적은 경우
        if not v: v = f'{cid}_{idx}'                             # 비어 있으면 카드아이디_순번 파일이 있는지 확인
        if '.' in v: return v
        for ext in ('webp', 'jpg', 'png', 'mp4', 'webm'):
            if os.path.exists(os.path.join(fxdir, f'{v}.{ext}')): return f'{v}.{ext}'
        return ''
    ws = wb['스킬']
    for r in list(ws.iter_rows(values_only=True))[1:]:
        if not r or not r[0] or not r[2]: continue
        cid = str(r[0]).strip()
        idx = len(skills.get(cid, [])) + 1
        skills.setdefault(cid, []).append({
            'card': str(r[1] or '').strip(), 'name': str(r[2]).strip(),
            'type': str(r[3] or '').strip(), 'mode': 'passive' if '패시브' in str(r[4]) else 'active',
            'target': str(r[5] or '').strip(), 'desc': str(r[6] or '').strip(), 'value': '' if r[7] is None else str(r[7]),
            'duration': num(r[8]), 'mp': num(r[9]), 'cooldown': num(r[10]), 'unlock': num(r[11], 1) or 1,
            'fx': fxname(cid, idx, r[12]), 'note': str(r[13] or '').strip(),
        })
    combos = []
    ws = wb['특정 조합']
    for r in list(ws.iter_rows(values_only=True))[1:]:
        if not r or not r[0]: continue
        eff = str(r[6] or '')
        combos.append({
            'name': str(r[0]).strip(), 'cards': [str(x).strip() for x in r[1:5] if x],
            'scope': 'all' if '전체' in str(r[5]) else 'combo',
            'atk': '공' in eff, 'def': '방' in eff, 'hp': '체' in eff,
            'value': float(r[7] or 0), 'desc': eff,
        })
    json.dump(skills, open(os.path.join(root, 'skills.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    json.dump(combos, open(os.path.join(root, 'combos.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    print(f'skills: {sum(len(v) for v in skills.values())} ({len(skills)} cards), combos: {len(combos)}')

if __name__ == '__main__':
    main(sys.argv[1])
